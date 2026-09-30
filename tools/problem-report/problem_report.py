import argparse
import json
import sys
import urllib.error
import urllib.request
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from urllib.parse import urlencode, urlparse
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PAGE_SIZE = 1000
MAX_LISTED_VALUES = 5
MAX_LISTED_LENGTH = 16

LABELS = {
    "missing-resource": "Finnes ikke",
    "missing-back-link-adapter": "Mangler lenke tilbake",
    "missing-back-link-autorelation": "Mangler lenke tilbake (FINT)",
    "unknown-link": "Ugyldig lenke",
}

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
RESOURCE_FILL = PatternFill("solid", fgColor="D9E1F2")
BOLD = Font(bold=True)
TITLE_FONT = Font(bold=True, size=16)
WRAP = Alignment(wrap_text=True, vertical="top")
NUMBER = "#,##0"


@dataclass
class Problem:
    component: str
    resource: str
    relation: str
    problem_type: str
    inverse: str
    source: str
    target: str


@dataclass
class Group:
    component: str
    resource: str
    relation: str
    problem_type: str
    inverse: str
    target_name: str
    target_component: str
    count: int = 0
    target_values: set = field(default_factory=set)
    example: tuple = None


def get_json(url, attempts=3):
    for attempt in range(1, attempts + 1):
        try:
            with urllib.request.urlopen(url, timeout=300) as response:
                return json.load(response)
        except urllib.error.HTTPError as error:
            if error.code == 404:
                return None
            raise
        except TimeoutError:
            if attempt == attempts:
                raise
            print(f"  timeout, retrying {url}", file=sys.stderr)


def fetch_problems(base_url, org, summary):
    for component in summary["summary"]["components"]:
        for resource in component["resources"]:
            for problem_type, count in resource["byProblemType"].items():
                if count > 0:
                    yield from fetch_problem_pages(
                        base_url, org,
                        {"component": component["component"], "resource": resource["resource"], "problemType": problem_type},
                    )


def fetch_problem_pages(base_url, org, filters):
    page = 0
    while True:
        query = urlencode({**filters, "page": page, "size": PAGE_SIZE})
        body = get_json(f"{base_url}/report/{org}/problems?{query}")
        if body is None:
            return
        yield from body["problems"]
        page += 1
        if page >= body["totalPages"]:
            return


def path_segments(href):
    return [s for s in urlparse(href).path.split("/") if s]


def target_component(href):
    segments = path_segments(href)
    return f"{segments[0]}_{segments[1]}" if len(segments) >= 5 else None


def target_name(href):
    segments = path_segments(href)
    return segments[-3] if len(segments) >= 5 else "?"


def record_id(href):
    segments = path_segments(href)
    return f"{segments[-2]}={segments[-1]}" if len(segments) >= 5 else href


def capitalize(text):
    return text[:1].upper() + text[1:]


def resource_title(component, resource):
    return f"{capitalize(resource)} ({component})"


def norwegian_number(value):
    return f"{value:,}".replace(",", " ")


def explanation(group, resource_total_refs, resources_with_records):
    if group.problem_type == "missing-resource":
        if (group.target_component, group.target_name) not in resources_with_records:
            return f"{capitalize(group.target_name)} har 0 poster i FINT, så alle lenkene hit feiler"
        text = f"Peker til {group.target_name} som ikke finnes"
        short = all(len(v) <= MAX_LISTED_LENGTH for v in group.target_values)
        if 0 < len(group.target_values) <= MAX_LISTED_VALUES and short:
            text += f" (verdier: {', '.join(sorted(group.target_values))})"
        if group.count == resource_total_refs:
            text += ". Alle lenkene feiler"
        return text
    if group.problem_type == "missing-back-link-adapter":
        return f"{capitalize(group.target_name)} peker ikke tilbake via '{group.inverse}'"
    if group.problem_type == "missing-back-link-autorelation":
        return (
            f"{capitalize(group.target_name)} peker ikke tilbake via '{group.inverse}'. "
            "Denne lenken skal FINT legge til, ikke adapteret"
        )
    return "Lenken har ugyldig format"


def resources_of(summary):
    return [r for c in summary["summary"]["components"] for r in c["resources"]]


def not_covered_from_summary(summary):
    counts = defaultdict(int)
    for resource in resources_of(summary):
        for group in resource.get("links", []):
            if group["scope"] == "not_covered":
                counts[group["targetComponent"]] += group["links"]
    return counts


def collect(org, base_url, summary):
    has_links = any(r.get("links") for r in resources_of(summary))
    scanned = set(summary["components"])
    problems = []
    not_scanned = not_covered_from_summary(summary) if has_links else defaultdict(int)
    for raw in fetch_problems(base_url, org, summary):
        if raw["problemType"] == "missing-resource" and not has_links:
            component = target_component(raw["targetHref"])
            if component not in scanned:
                not_scanned[component or "?"] += 1
                continue
        problems.append(Problem(
            component=raw["component"],
            resource=raw["resource"],
            relation=raw.get("relationName") or "-",
            problem_type=raw["problemType"],
            inverse=raw.get("expectedInverseName") or "-",
            source=raw["sourceSelf"],
            target=raw["targetHref"],
        ))
    return problems, not_scanned


def group_problems(problems):
    groups = {}
    for problem in problems:
        key = (
            problem.component, problem.resource, problem.relation, problem.problem_type,
            problem.inverse, target_name(problem.target), target_component(problem.target),
        )
        group = groups.setdefault(key, Group(*key))
        group.count += 1
        if group.example is None:
            group.example = (record_id(problem.source), record_id(problem.target))
        value = path_segments(problem.target)[-1] if path_segments(problem.target) else ""
        if "*" not in value and len(group.target_values) <= MAX_LISTED_VALUES:
            group.target_values.add(value)
    return list(groups.values())


def resource_order(groups):
    totals = defaultdict(int)
    for group in groups:
        totals[(group.component, group.resource)] += group.count
    return totals


def write_header_row(sheet, row, headers):
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=column, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT


def set_widths(sheet, widths):
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def scanned_at(summary):
    utc = datetime.strptime(summary["scanCompletedAt"][:19], "%Y-%m-%dT%H:%M:%S").replace(tzinfo=timezone.utc)
    return utc.astimezone(ZoneInfo("Europe/Oslo"))


def write_overview(workbook, org, env, summary, groups, not_scanned):
    sheet = workbook.active
    sheet.title = "Oversikt"
    set_widths(sheet, [34, 26, 10, 60, 44, 50])

    has_links = any(r.get("links") for r in resources_of(summary))
    total_refs = summary["summary"]["totalRefs"] - (0 if has_links else sum(not_scanned.values()))
    total_broken = sum(g.count for g in groups)

    sheet.cell(row=1, column=1, value=f"Lenkefeil i FINT ({env})").font = TITLE_FONT
    facts = [
        ("Organisasjon", org, None),
        ("Skannet", scanned_at(summary).strftime("%d.%m.%Y %H:%M"), None),
        ("Lenker sjekket", total_refs, NUMBER),
        ("Lenker med feil", total_broken, NUMBER),
        ("Andel OK", 1 - total_broken / total_refs if total_refs > 0 else None, "0.00%"),
    ]
    for row, (label, value, number_format) in enumerate(facts, start=3):
        sheet.cell(row=row, column=1, value=label).font = BOLD
        cell = sheet.cell(row=row, column=2, value=value)
        cell.alignment = Alignment(horizontal="left")
        if number_format:
            cell.number_format = number_format

    row = 3 + len(facts) + 1
    write_header_row(sheet, row, ["Ressurs / relasjon", "Feil", "Antall", "Forklaring", "Eksempel: fra", "Eksempel: til"])
    sheet.freeze_panes = sheet.cell(row=row + 1, column=1)

    resources_with_records = {
        (c["component"], r["resource"])
        for c in summary["summary"]["components"] for r in c["resources"] if r["totalRecords"] > 0
    }
    total_refs_by_resource = {
        (c["component"], r["resource"]): r["totalRefs"]
        for c in summary["summary"]["components"] for r in c["resources"]
    }
    totals = resource_order(groups)
    by_resource = defaultdict(list)
    for group in groups:
        by_resource[(group.component, group.resource)].append(group)

    for key in sorted(by_resource, key=lambda k: -totals[k]):
        row += 1
        sheet.cell(row=row, column=1, value=resource_title(*key))
        sheet.cell(row=row, column=3, value=totals[key]).number_format = NUMBER
        for column in range(1, 7):
            cell = sheet.cell(row=row, column=column)
            cell.fill = RESOURCE_FILL
            cell.font = BOLD
        for group in sorted(by_resource[key], key=lambda g: -g.count):
            row += 1
            sheet.cell(row=row, column=1, value=group.relation).alignment = Alignment(indent=2)
            sheet.cell(row=row, column=2, value=LABELS.get(group.problem_type, group.problem_type))
            sheet.cell(row=row, column=3, value=group.count).number_format = NUMBER
            sheet.cell(row=row, column=4, value=explanation(group, total_refs_by_resource.get(key), resources_with_records)).alignment = WRAP
            sheet.cell(row=row, column=5, value=group.example[0]).alignment = WRAP
            sheet.cell(row=row, column=6, value=group.example[1]).alignment = WRAP

    if not_scanned:
        row += 2
        note = (
            f"Ikke sjekket: {norwegian_number(sum(not_scanned.values()))} lenker peker til "
            f"{', '.join(sorted(not_scanned))}. Skanningen hentet ikke disse dataene, "
            "så vi vet ikke om lenkene er riktige. De er ikke med i tallene over."
        )
        cell = sheet.cell(row=row, column=1, value=note)
        cell.font = Font(italic=True)
        cell.alignment = WRAP
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        sheet.row_dimensions[row].height = 32


def write_all_errors(workbook, problems, groups):
    sheet = workbook.create_sheet("Alle feil")
    write_header_row(sheet, 1, ["Ressurs", "Relasjon", "Feil", "Fra (id)", "Til (id)", "Fra (lenke)", "Til (lenke)"])
    set_widths(sheet, [34, 24, 26, 44, 44, 80, 80])
    sheet.freeze_panes = "A2"

    totals = resource_order(groups)
    group_counts = {(g.component, g.resource, g.relation, g.problem_type): g.count for g in groups}

    def sort_key(problem):
        return (
            -totals[(problem.component, problem.resource)], problem.component, problem.resource,
            -group_counts[(problem.component, problem.resource, problem.relation, problem.problem_type)],
            problem.relation, problem.problem_type,
        )

    for problem in sorted(problems, key=sort_key):
        sheet.append([
            resource_title(problem.component, problem.resource),
            problem.relation,
            LABELS.get(problem.problem_type, problem.problem_type),
            record_id(problem.source),
            record_id(problem.target),
            problem.source,
            problem.target,
        ])
    if sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions


def main():
    parser = argparse.ArgumentParser(description="Write the latest link-walker errors for one org to an Excel report.")
    parser.add_argument("--org", required=True, help="Org id, e.g. ude_oslo_kommune_no")
    parser.add_argument("--env", default="beta")
    parser.add_argument("--base-url", default="http://localhost:18080/link-walker")
    parser.add_argument("--out")
    args = parser.parse_args()

    summary = get_json(f"{args.base_url}/report/{args.org}/summary")
    if summary is None:
        sys.exit(f"{args.org}: no scan found")

    problems, not_scanned = collect(args.org, args.base_url, summary)
    groups = group_problems(problems)
    out = args.out or f"lenkefeil-{args.org}-{args.env}-{datetime.now():%Y-%m-%d}.xlsx"

    workbook = Workbook()
    write_overview(workbook, args.org, args.env, summary, groups, not_scanned)
    write_all_errors(workbook, problems, groups)
    workbook.save(out)
    print(
        f"{args.org}: {len(problems)} errors in {len(groups)} relations, "
        f"{sum(not_scanned.values())} not checked",
        file=sys.stderr,
    )
    print(f"Wrote {out}", file=sys.stderr)


if __name__ == "__main__":
    main()
